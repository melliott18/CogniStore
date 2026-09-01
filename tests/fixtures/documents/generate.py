"""Regenerate the deterministic binary fixtures for ticket #33.

The generated content is intentionally small and project-authored.  Keep the
visible text and metadata constants in sync with README.md when changing this
file.
"""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, ByteStringObject
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

VISIBLE_LINES = (
    "CogniStore Document Fixture",
    "This project-authored sample verifies deterministic text extraction.",
    "Alpha records stay searchable. Beta pages remain reproducible.",
)
EXPECTED_NORMALIZED_TEXT = "\n".join(VISIBLE_LINES)

TITLE = "CogniStore Document Fixture"
AUTHOR = "CogniStore Contributors"
SUBJECT = "Licensed deterministic extraction fixture"
KEYWORDS = "cognistore, extraction, fixture"
GENERATOR_NAME = "CogniStore fixture generator"
FIXED_TIMESTAMP = datetime(2000, 1, 1, tzinfo=timezone.utc)
ENCRYPTED_PDF_PASSWORD = "fixture-password"
ENCRYPTED_PDF_OWNER_PASSWORD = "fixture-owner-password"
ENCRYPTED_PDF_FILE_ID = b"CogniStoreFixture"

ZIP_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)


def _write_zip(path: Path, entries: dict[str, bytes]) -> None:
    """Write *entries* with stable ordering, timestamps, permissions, and compression."""

    with ZipFile(path, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(entries):
            info = ZipInfo(name, date_time=ZIP_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(info, entries[name], compress_type=ZIP_DEFLATED, compresslevel=9)


def _normalize_docx_package(source: Path, destination: Path) -> None:
    with ZipFile(source, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    _write_zip(destination, entries)


def _set_style_font(style: object, name: str) -> None:
    style.font.name = name  # type: ignore[attr-defined]
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()  # type: ignore[attr-defined]
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)


def generate_sample_pdf(path: Path) -> None:
    buffer = BytesIO()
    document = canvas.Canvas(
        buffer,
        pagesize=LETTER,
        pageCompression=1,
        invariant=1,
        pdfVersion=(1, 4),
    )
    document.setTitle(TITLE)
    document.setAuthor(AUTHOR)
    document.setSubject(SUBJECT)
    document.setKeywords(KEYWORDS)
    document.setCreator(GENERATOR_NAME)
    document._doc.info.producer = GENERATOR_NAME

    page_width, page_height = LETTER
    del page_width
    document.setFont("Helvetica-Bold", 16)
    document.drawString(inch, page_height - inch, VISIBLE_LINES[0])
    document.setFont("Helvetica", 11)
    document.drawString(inch, page_height - 1.45 * inch, VISIBLE_LINES[1])
    document.drawString(inch, page_height - 1.75 * inch, VISIBLE_LINES[2])
    document.showPage()
    document.save()
    path.write_bytes(buffer.getvalue())


def generate_encrypted_pdf(source: Path, destination: Path) -> None:
    reader = PdfReader(source)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    fixed_file_id = ByteStringObject(ENCRYPTED_PDF_FILE_ID)
    writer._ID = ArrayObject((fixed_file_id, fixed_file_id))
    writer.encrypt(
        user_password=ENCRYPTED_PDF_PASSWORD,
        owner_password=ENCRYPTED_PDF_OWNER_PASSWORD,
        algorithm="RC4-128",
    )
    with destination.open("wb") as stream:
        writer.write(stream)


def generate_sample_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading = document.styles["Heading 1"]
    _set_style_font(heading, "Calibri")
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.line_spacing = 1.1

    document.add_heading(VISIBLE_LINES[0], level=1)
    document.add_paragraph(VISIBLE_LINES[1])
    document.add_paragraph(VISIBLE_LINES[2])

    properties = document.core_properties
    properties.title = TITLE
    properties.author = AUTHOR
    properties.subject = SUBJECT
    properties.keywords = KEYWORDS
    properties.category = "Test Fixture"
    properties.comments = "Project-authored test fixture."
    properties.content_status = "Final"
    properties.identifier = "cognistore-ticket-33-fixture"
    properties.language = "en-US"
    properties.last_modified_by = GENERATOR_NAME
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP
    properties.revision = 1
    properties.version = "1.0"

    with TemporaryDirectory(prefix="cognistore-docx-") as temp_dir:
        intermediate = Path(temp_dir) / "sample.docx"
        document.save(intermediate)
        _normalize_docx_package(intermediate, path)


def generate_corrupt_pdf(path: Path) -> None:
    path.write_bytes(b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog\n")


def generate_corrupt_docx(path: Path) -> None:
    content_types = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="{DOCX_CONTENT_TYPE}"/>
</Types>
""".encode()
    relationships = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""
    malformed_document = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body><w:p><w:r><w:t>Intentionally truncated fixture
"""
    _write_zip(
        path,
        {
            "[Content_Types].xml": content_types,
            "_rels/.rels": relationships,
            "word/document.xml": malformed_document,
        },
    )


def generate_all(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    sample_pdf = output_dir / "sample.pdf"
    generate_sample_pdf(sample_pdf)
    generate_encrypted_pdf(sample_pdf, output_dir / "encrypted.pdf")
    generate_sample_docx(output_dir / "sample.docx")
    generate_corrupt_pdf(output_dir / "corrupt.pdf")
    generate_corrupt_docx(output_dir / "corrupt.docx")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="directory for generated binary fixtures (default: this script's directory)",
    )
    args = parser.parse_args()
    generate_all(args.output_dir.resolve())


if __name__ == "__main__":
    main()
