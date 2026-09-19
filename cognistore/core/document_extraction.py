from __future__ import annotations

import json
import logging
import math
import multiprocessing
import os
import pickle
import re
import sys
import threading
import unicodedata
from collections.abc import Mapping, Sequence
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timezone
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO
from multiprocessing.connection import Connection
from multiprocessing.shared_memory import SharedMemory
from time import monotonic
from typing import Any, Literal, Protocol

PDF_MIME_TYPE = "application/pdf"
DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
SUPPORTED_DOCUMENT_MIME_TYPES = frozenset({PDF_MIME_TYPE, DOCX_MIME_TYPE})

DEFAULT_MAX_FILE_BYTES = 25 * 1024 * 1024
DEFAULT_TIMEOUT_SECONDS = 30.0
DEFAULT_MAX_OUTPUT_BYTES = 4 * 1024 * 1024
EXTRACTION_SCHEMA_VERSION = 1
NORMALIZATION_VERSION = 1

ExtractionStatus = Literal["succeeded", "failed"]
ExtractionFailureCode = Literal[
    "unsupported_mime",
    "input_too_large",
    "timeout",
    "encrypted",
    "corrupt",
    "output_too_large",
    "parser_error",
    "worker_error",
]

_HORIZONTAL_WHITESPACE = re.compile(r"[^\S\n]+")
_OLE_COMPOUND_FILE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")


@dataclass(frozen=True)
class ExtractionLimits:
    """Resource bounds applied to every document extraction attempt."""

    max_file_bytes: int = DEFAULT_MAX_FILE_BYTES
    timeout_seconds: float = DEFAULT_TIMEOUT_SECONDS
    max_output_bytes: int = DEFAULT_MAX_OUTPUT_BYTES

    def __post_init__(self) -> None:
        for name in ("max_file_bytes", "max_output_bytes"):
            value = getattr(self, name)
            if (
                isinstance(value, bool)
                or not isinstance(value, int)
                or value <= 0
                or value > sys.maxsize
            ):
                raise ValueError(
                    f"{name} must be a positive integer no greater than sys.maxsize"
                )
        if (
            isinstance(self.timeout_seconds, bool)
            or not isinstance(self.timeout_seconds, (int, float))
            or not math.isfinite(self.timeout_seconds)
            or self.timeout_seconds <= 0
            or self.timeout_seconds > threading.TIMEOUT_MAX
        ):
            raise ValueError(
                "timeout_seconds must be positive, finite, and no greater than "
                "threading.TIMEOUT_MAX"
            )


@dataclass(frozen=True)
class ParserIdentity:
    name: str
    implementation_version: str
    runtime_version: str

    def to_metadata(self) -> dict[str, str]:
        return {
            "name": self.name,
            "implementation_version": self.implementation_version,
            "runtime_version": self.runtime_version,
        }


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    metadata: Mapping[str, object]


class DocumentParserAdapter(Protocol):
    """One pickle-compatible parser implementation used by an isolated worker."""

    identity: ParserIdentity
    mime_types: frozenset[str]

    def parse(self, data: bytes, *, mime: str) -> ParsedDocument: ...


class DocumentParseError(RuntimeError):
    """A safely classifiable parser failure with no document content attached."""

    def __init__(self, code: ExtractionFailureCode) -> None:
        if code not in {"encrypted", "corrupt"}:
            raise ValueError("parser failures must be encrypted or corrupt")
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class DocumentExtraction:
    status: ExtractionStatus
    source_mime: str | None
    source_size: int
    parser: ParserIdentity | None
    text: str | None = None
    text_bytes: int = 0
    output_bytes: int = 0
    document_metadata: Mapping[str, object] | None = None
    failure_code: ExtractionFailureCode | None = None

    def to_metadata(self) -> dict[str, object]:
        """Return the stable, JSON-safe catalog representation."""

        return {
            "schema_version": EXTRACTION_SCHEMA_VERSION,
            "status": self.status,
            "source_mime": self.source_mime,
            "source_size": self.source_size,
            "parser": self.parser.to_metadata() if self.parser is not None else None,
            "normalization_version": NORMALIZATION_VERSION,
            "text": self.text,
            "text_bytes": self.text_bytes,
            "output_bytes": self.output_bytes,
            "document_metadata": deepcopy(dict(self.document_metadata or {})),
            "failure_code": self.failure_code,
        }


@dataclass(frozen=True)
class _WorkerOutcome:
    text: str | None = None
    text_bytes: int = 0
    output_bytes: int = 0
    metadata: Mapping[str, object] | None = None
    failure_code: ExtractionFailureCode | None = None


@dataclass
class _WorkerReceiveState:
    outcome: _WorkerOutcome | None = None
    failed: bool = False


@dataclass(frozen=True)
class _RegisteredParser:
    identity: ParserIdentity
    serialized_adapter: bytes
    serialized_mime: bytes


class PdfParserAdapter:
    """Extract PDF text and document-information metadata with pypdf."""

    mime_types = frozenset({PDF_MIME_TYPE})

    def __init__(self) -> None:
        self.identity = ParserIdentity(
            name="pypdf",
            implementation_version="1",
            runtime_version=_package_version("pypdf"),
        )

    def parse(self, data: bytes, *, mime: str) -> ParsedDocument:
        if mime not in self.mime_types:
            raise ValueError(f"unsupported PDF MIME type: {mime}")
        try:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise DocumentParseError("encrypted")
            pages = [page.extract_text() or "" for page in reader.pages]
            information = reader.metadata
            metadata: dict[str, object] = {
                "format": "pdf",
                "title": _metadata_text(_safe_attribute(information, "title")),
                "author": _metadata_text(_safe_attribute(information, "author")),
                "subject": _metadata_text(_safe_attribute(information, "subject")),
                "keywords": _metadata_text(_safe_attribute(information, "keywords")),
                "language": None,
                "created_at": _metadata_datetime(
                    _safe_attribute(information, "creation_date")
                ),
                "modified_at": _metadata_datetime(
                    _safe_attribute(information, "modification_date")
                ),
                "page_count": len(reader.pages),
            }
            return ParsedDocument(text="\n\n".join(pages), metadata=metadata)
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("corrupt") from exc


class DocxParserAdapter:
    """Extract DOCX body blocks and core properties with python-docx."""

    mime_types = frozenset({DOCX_MIME_TYPE})

    def __init__(self) -> None:
        self.identity = ParserIdentity(
            name="python-docx",
            implementation_version="1",
            runtime_version=_package_version("python-docx"),
        )

    def parse(self, data: bytes, *, mime: str) -> ParsedDocument:
        if mime not in self.mime_types:
            raise ValueError(f"unsupported DOCX MIME type: {mime}")
        if data.startswith(_OLE_COMPOUND_FILE_SIGNATURE):
            # Password-protected OOXML packages use an OLE encrypted container,
            # not the ZIP package consumed by python-docx.
            raise DocumentParseError("encrypted")
        try:
            from docx import Document
            from docx.table import Table

            document = Document(BytesIO(data))
            blocks: list[str] = []
            for block in document.iter_inner_content():
                if isinstance(block, Table):
                    for row in block.rows:
                        blocks.append(" | ".join(cell.text for cell in row.cells))
                else:
                    blocks.append(block.text)

            properties = document.core_properties
            metadata: dict[str, object] = {
                "format": "docx",
                "title": _metadata_text(_safe_attribute(properties, "title")),
                "author": _metadata_text(_safe_attribute(properties, "author")),
                "subject": _metadata_text(_safe_attribute(properties, "subject")),
                "keywords": _metadata_text(_safe_attribute(properties, "keywords")),
                "language": _metadata_text(_safe_attribute(properties, "language")),
                "created_at": _metadata_datetime(
                    _safe_attribute(properties, "created")
                ),
                "modified_at": _metadata_datetime(
                    _safe_attribute(properties, "modified")
                ),
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            }
            return ParsedDocument(text="\n".join(blocks), metadata=metadata)
        except DocumentParseError:
            raise
        except Exception as exc:
            raise DocumentParseError("corrupt") from exc


class DocumentExtractionPipeline:
    """Select a parser, enforce bounds, and return a non-raising outcome."""

    def __init__(
        self,
        parsers: Sequence[DocumentParserAdapter] | None = None,
        *,
        limits: ExtractionLimits | None = None,
        process_start_method: str = "spawn",
    ) -> None:
        configured = tuple(parsers) if parsers is not None else (
            PdfParserAdapter(),
            DocxParserAdapter(),
        )
        parser_by_mime: dict[str, _RegisteredParser] = {}
        for parser in configured:
            if not parser.mime_types:
                raise ValueError("document parser must declare at least one MIME type")
            try:
                serialized_adapter = pickle.dumps(
                    parser,
                    protocol=pickle.HIGHEST_PROTOCOL,
                )
            except Exception as exc:
                raise ValueError("document parser must be pickle-compatible") from exc
            for mime in parser.mime_types:
                if not isinstance(mime, str) or not mime:
                    raise ValueError(
                        "document parser MIME types must be non-empty strings"
                    )
                try:
                    serialized_mime = mime.encode("utf-8")
                except UnicodeEncodeError as exc:
                    raise ValueError(
                        "document parser MIME types must be valid UTF-8"
                    ) from exc
                if mime in parser_by_mime:
                    raise ValueError(f"multiple document parsers registered for {mime}")
                parser_by_mime[mime] = _RegisteredParser(
                    identity=parser.identity,
                    serialized_adapter=serialized_adapter,
                    serialized_mime=serialized_mime,
                )
        self._parsers = parser_by_mime
        self.limits = limits or ExtractionLimits()
        self._process_start_method = process_start_method

    def needs_document_bytes(self, mime: str | None, *, source_size: int) -> bool:
        """Return whether a caller should load the complete bounded object."""

        return (
            mime in self._parsers
            and _valid_source_size(source_size)
            and source_size <= self.limits.max_file_bytes
        )

    def extract(
        self,
        data: bytes,
        *,
        mime: str | None,
        source_size: int | None = None,
    ) -> DocumentExtraction:
        if source_size is None:
            source_size = len(data)
        if not _valid_source_size(source_size):
            raise ValueError("source_size must be a non-negative integer")

        registration = self._parsers.get(mime or "")
        if registration is None:
            return self._failure(
                mime=mime,
                source_size=source_size,
                parser=None,
                code="unsupported_mime",
            )
        if source_size > self.limits.max_file_bytes or len(data) > self.limits.max_file_bytes:
            return self._failure(
                mime=mime,
                source_size=source_size,
                parser=registration.identity,
                code="input_too_large",
            )
        if len(data) != source_size:
            # Supported, in-limit documents must be complete. Scan generation
            # checks will discard concurrent mutations, while direct callers
            # still receive a safe non-success outcome for partial input.
            return self._failure(
                mime=mime,
                source_size=source_size,
                parser=registration.identity,
                code="worker_error",
            )

        assert mime is not None
        outcome = self._run_parser(registration, data)
        if outcome.failure_code is not None:
            return self._failure(
                mime=mime,
                source_size=source_size,
                parser=registration.identity,
                code=outcome.failure_code,
            )
        return DocumentExtraction(
            status="succeeded",
            source_mime=mime,
            source_size=source_size,
            parser=registration.identity,
            text=outcome.text or "",
            text_bytes=outcome.text_bytes,
            output_bytes=outcome.output_bytes,
            document_metadata=dict(outcome.metadata or {}),
            failure_code=None,
        )

    def _run_parser(
        self,
        registration: _RegisteredParser,
        data: bytes,
    ) -> _WorkerOutcome:
        receive_connection: Connection | None = None
        send_connection: Connection | None = None
        process: multiprocessing.Process | None = None
        receiver_thread: threading.Thread | None = None
        shared_input: SharedMemory | None = None
        shared_adapter: SharedMemory | None = None
        deadline = monotonic() + self.limits.timeout_seconds
        try:
            context: Any = multiprocessing.get_context(self._process_start_method)
            shared_input = SharedMemory(create=True, size=max(1, len(data)))
            shared_input_buffer = shared_input.buf
            if shared_input_buffer is None:
                raise BufferError("shared input buffer is unavailable")
            try:
                shared_input_buffer[: len(data)] = data
            finally:
                shared_input_buffer.release()
            serialized_adapter = registration.serialized_adapter
            serialized_mime = registration.serialized_mime
            shared_adapter = SharedMemory(
                create=True,
                size=max(1, len(serialized_adapter) + len(serialized_mime)),
            )
            shared_adapter_buffer = shared_adapter.buf
            if shared_adapter_buffer is None:
                raise BufferError("shared adapter buffer is unavailable")
            try:
                shared_adapter_buffer[: len(serialized_adapter)] = serialized_adapter
                shared_adapter_buffer[
                    len(serialized_adapter) : len(serialized_adapter)
                    + len(serialized_mime)
                ] = serialized_mime
            finally:
                shared_adapter_buffer.release()
            receive_connection, send_connection = context.Pipe(duplex=False)
            process = context.Process(
                target=_parser_worker,
                args=(
                    send_connection,
                    shared_input.name,
                    len(data),
                    shared_adapter.name,
                    len(serialized_adapter),
                    len(serialized_mime),
                    self.limits.max_output_bytes,
                ),
                daemon=True,
            )
        except (OSError, ValueError, BufferError):
            if receive_connection is not None:
                _close_connection(receive_connection)
            if send_connection is not None:
                _close_connection(send_connection)
            if shared_input is not None:
                _release_shared_memory(shared_input)
            if shared_adapter is not None:
                _release_shared_memory(shared_adapter)
            return _WorkerOutcome(failure_code="worker_error")
        assert receive_connection is not None
        assert send_connection is not None
        assert process is not None
        try:
            if monotonic() >= deadline:
                return _WorkerOutcome(failure_code="timeout")
            try:
                process.start()
            except Exception:
                return _WorkerOutcome(failure_code="worker_error")
            finally:
                send_connection.close()

            remaining = deadline - monotonic()
            if remaining <= 0:
                _stop_process(process)
                return _WorkerOutcome(failure_code="timeout")

            receive_state = _WorkerReceiveState()
            receive_finished = threading.Event()
            receiver_thread = threading.Thread(
                target=_receive_worker_outcome,
                args=(receive_connection, receive_state, receive_finished),
                daemon=True,
                name="cognistore-document-parser-result",
            )
            try:
                receiver_thread.start()
            except RuntimeError:
                return _WorkerOutcome(failure_code="worker_error")
            remaining = deadline - monotonic()
            if remaining <= 0:
                _stop_process(process)
                return _WorkerOutcome(failure_code="timeout")
            if not receive_finished.wait(
                timeout=min(remaining, threading.TIMEOUT_MAX)
            ):
                _stop_process(process)
                return _WorkerOutcome(failure_code="timeout")
            if monotonic() >= deadline:
                _stop_process(process)
                return _WorkerOutcome(failure_code="timeout")
            if receive_state.failed or receive_state.outcome is None:
                return _WorkerOutcome(failure_code="worker_error")
            return receive_state.outcome
        finally:
            if process is not None and process.pid is not None:
                try:
                    process.join(timeout=0.2)
                    if process.is_alive():
                        _stop_process(process)
                    if not process.is_alive():
                        process.close()
                except (AssertionError, OSError, ValueError):
                    pass
            if receiver_thread is not None and receiver_thread.is_alive():
                receiver_thread.join(timeout=0.2)
            if receive_connection is not None:
                _close_connection(receive_connection)
            if send_connection is not None:
                _close_connection(send_connection)
            if shared_input is not None:
                _release_shared_memory(shared_input)
            if shared_adapter is not None:
                _release_shared_memory(shared_adapter)

    @staticmethod
    def _failure(
        *,
        mime: str | None,
        source_size: int,
        parser: ParserIdentity | None,
        code: ExtractionFailureCode,
    ) -> DocumentExtraction:
        return DocumentExtraction(
            status="failed",
            source_mime=mime,
            source_size=source_size,
            parser=parser,
            failure_code=code,
        )


def normalize_document_text(value: str) -> str:
    """Apply the version-one deterministic text normalization contract."""

    text = unicodedata.normalize("NFC", value)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\t", " ").replace("\v", " ").replace("\f", " ")
    text = "".join(
        character
        for character in text
        if character == "\n" or unicodedata.category(character) != "Cc"
    )

    lines = [_HORIZONTAL_WHITESPACE.sub(" ", line).strip() for line in text.split("\n")]
    normalized: list[str] = []
    previous_was_blank = False
    for line in lines:
        if line:
            normalized.append(line)
            previous_was_blank = False
        elif normalized and not previous_was_blank:
            normalized.append("")
            previous_was_blank = True
    while normalized and not normalized[-1]:
        normalized.pop()
    return "\n".join(normalized)


def _parser_worker(
    connection: Connection,
    shared_input_name: str,
    input_size: int,
    shared_adapter_name: str,
    adapter_size: int,
    mime_size: int,
    max_output_bytes: int,
) -> None:
    # Parser diagnostics can quote arbitrary document strings (for example,
    # malformed PDF dictionary keys). Silence Python and native output before
    # adapter unpickling; the parent receives only the structured outcome.
    with open(os.devnull, "w") as sink:
        os.dup2(sink.fileno(), 1)
        os.dup2(sink.fileno(), 2)
        sys.stdout = sink
        sys.stderr = sink
        logging.disable(sys.maxsize)
        _run_parser_worker(
            connection,
            shared_input_name,
            input_size,
            shared_adapter_name,
            adapter_size,
            mime_size,
            max_output_bytes,
        )


def _run_parser_worker(
    connection: Connection,
    shared_input_name: str,
    input_size: int,
    shared_adapter_name: str,
    adapter_size: int,
    mime_size: int,
    max_output_bytes: int,
) -> None:
    try:
        shared_input: SharedMemory | None = None
        shared_adapter: SharedMemory | None = None
        data: bytes | None = None
        parser: DocumentParserAdapter | None = None
        try:
            shared_input = SharedMemory(name=shared_input_name)
            shared_input_buffer = shared_input.buf
            if shared_input_buffer is None:
                raise BufferError("shared input buffer is unavailable")
            try:
                data = bytes(shared_input_buffer[:input_size])
            finally:
                shared_input_buffer.release()
            shared_adapter = SharedMemory(name=shared_adapter_name)
            shared_adapter_buffer = shared_adapter.buf
            if shared_adapter_buffer is None:
                raise BufferError("shared adapter buffer is unavailable")
            try:
                serialized_adapter = bytes(shared_adapter_buffer[:adapter_size])
                mime = bytes(
                    shared_adapter_buffer[adapter_size : adapter_size + mime_size]
                ).decode("utf-8")
            finally:
                shared_adapter_buffer.release()
            # Adapter bytes are produced from trusted in-process configuration,
            # never from an indexed object or other external input.
            parser = pickle.loads(serialized_adapter)  # nosec B301
        except Exception:
            outcome = _WorkerOutcome(failure_code="worker_error")
        finally:
            if shared_input is not None:
                shared_input.close()
            if shared_adapter is not None:
                shared_adapter.close()

        if data is not None and parser is not None:
            try:
                parsed = parser.parse(data, mime=mime)
                text = normalize_document_text(parsed.text)
                metadata_json = json.dumps(
                    parsed.metadata,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                    allow_nan=False,
                )
                metadata = json.loads(metadata_json)
                if not isinstance(metadata, dict):
                    raise TypeError("document metadata must be a JSON object")
                text_bytes = len(text.encode("utf-8"))
                output_bytes = text_bytes + len(metadata_json.encode("utf-8"))
                if output_bytes > max_output_bytes:
                    outcome = _WorkerOutcome(failure_code="output_too_large")
                else:
                    outcome = _WorkerOutcome(
                        text=text,
                        text_bytes=text_bytes,
                        output_bytes=output_bytes,
                        metadata=metadata,
                    )
            except DocumentParseError as exc:
                outcome = _WorkerOutcome(failure_code=exc.code)
            except Exception:
                outcome = _WorkerOutcome(failure_code="parser_error")
        connection.send(outcome)
    except (BrokenPipeError, EOFError, OSError):
        pass
    finally:
        connection.close()


def _receive_worker_outcome(
    connection: Connection,
    state: _WorkerReceiveState,
    finished: threading.Event,
) -> None:
    """Receive one framed pipe message without blocking the deadline owner."""

    try:
        outcome = connection.recv()
        if isinstance(outcome, _WorkerOutcome):
            state.outcome = outcome
        else:
            state.failed = True
    except Exception:
        state.failed = True
    finally:
        finished.set()


def _stop_process(process: multiprocessing.Process) -> None:
    try:
        if not process.is_alive():
            process.join(timeout=0)
            return
        process.terminate()
        process.join(timeout=0.2)
        if process.is_alive() and hasattr(process, "kill"):
            process.kill()
            process.join(timeout=0.2)
    except (OSError, ValueError):
        # Cleanup must not replace the stable per-object timeout/worker result.
        return


def _release_shared_memory(shared_input: SharedMemory) -> None:
    try:
        shared_input.close()
    except (BufferError, OSError):
        pass
    try:
        shared_input.unlink()
    except OSError:
        pass


def _close_connection(connection: Connection) -> None:
    try:
        connection.close()
    except (OSError, ValueError):
        pass


def _valid_source_size(value: object) -> bool:
    return isinstance(value, int) and not isinstance(value, bool) and value >= 0


def _package_version(package: str) -> str:
    try:
        return version(package)
    except PackageNotFoundError:
        return "unavailable"


def _safe_attribute(value: object, name: str) -> object:
    if value is None:
        return None
    try:
        return getattr(value, name, None)
    except Exception:
        return None


def _metadata_text(value: object) -> str | None:
    if value is None:
        return None
    text = normalize_document_text(str(value)).replace("\n", " ")
    return text or None


def _metadata_datetime(value: object) -> str | None:
    if not isinstance(value, datetime):
        return _metadata_text(value)
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    else:
        value = value.astimezone(timezone.utc)
    return value.isoformat(timespec="seconds").replace("+00:00", "Z")
