"""Generate the deterministic, sanitized content-search sample corpus.

The corpus documents are original CogniStore project material.  This script is
for maintainers; ordinary sample loading uses the checked-in files after
verifying ``SHA256SUMS``.
"""

from __future__ import annotations

import argparse
import filecmp
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from importlib.metadata import version
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor  # type: ignore[import-untyped,unused-ignore]
from reportlab.lib.pagesizes import LETTER  # type: ignore[import-untyped,unused-ignore]
from reportlab.lib.units import inch  # type: ignore[import-untyped,unused-ignore]
from reportlab.pdfbase.pdfmetrics import stringWidth  # type: ignore[import-untyped,unused-ignore]
from reportlab.pdfgen import canvas  # type: ignore[import-untyped,unused-ignore]

SCHEMA_VERSION = 1
CORPUS_NAME = "cognistore-content-search-sample-v1"
CORPUS_BUCKET = "sample-documents"
LICENSE_NAME = "MIT"

AUTHOR = "CogniStore Contributors"
GENERATOR_NAME = "CogniStore sample corpus generator"
FIXED_TIMESTAMP = datetime(2000, 1, 1, tzinfo=timezone.utc)
ZIP_TIMESTAMP = (2000, 1, 1, 0, 0, 0)

NAVY = HexColor("#183B56")
BLUE = HexColor("#2E74B5")
MUTED = HexColor("#5F6B76")
LIGHT_BLUE = HexColor("#E8F1F8")
INK = HexColor("#182026")

EXPECTED_DEPENDENCIES = {
    "python-docx": "1.2.0",
    "reportlab": "4.4.9",
}
GENERATED_FILES = (
    "database-backups.docx",
    "incident-response.pdf",
    "manifest.json",
    "records-retention.pdf",
    "SHA256SUMS",
)

_CONTENT_TYPES = "[Content_Types].xml"
_PACKAGE_RELATIONSHIPS = "_rels/.rels"
_DOCUMENT_RELATIONSHIPS = "word/_rels/document.xml.rels"
_REMOVED_DOCX_PARTS = (
    "customXml/_rels/item1.xml.rels",
    "customXml/item1.xml",
    "customXml/itemProps1.xml",
    "docProps/app.xml",
    "docProps/thumbnail.jpeg",
)
_CUSTOM_XML_RELATIONSHIP = (
    b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
)
_EXTENDED_PROPERTIES_RELATIONSHIP = (
    b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    b"extended-properties"
)
_THUMBNAIL_RELATIONSHIP = (
    b"http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail"
)
_RSID_ATTRIBUTE = re.compile(rb' w:rsid[A-Za-z]*="[0-9A-Fa-f]{8}"')
_RSID_LIST = re.compile(rb"<w:rsids>.*?</w:rsids>", re.DOTALL)
_RSID_ELEMENT = re.compile(rb"<w:rsid\b[^>]*/>")
_WORD_2010_DOCUMENT_ID = re.compile(rb"<w14:docId\b[^>]*/>")
_SAVE_PREVIEW_PICTURE = re.compile(rb"<w:savePreviewPicture\s*/>")


@dataclass(frozen=True)
class Section:
    heading: str
    paragraphs: tuple[str, ...]


@dataclass(frozen=True)
class DocumentSpec:
    filename: str
    title: str
    subtitle: str
    subject: str
    keywords: str
    purpose: str
    sections: tuple[Section, ...]


RETENTION = DocumentSpec(
    filename="records-retention.pdf",
    title="Records Retention and Legal Holds",
    subtitle="Sample policy - records governance",
    subject="Fictional records governance policy",
    keywords="records, retention, contracts, legal hold, deletion",
    purpose=(
        "This fictional policy demonstrates searchable retention rules without "
        "using customer, employee, or confidential data."
    ),
    sections=(
        Section(
            "Storage schedule",
            (
                "Active customer agreements remain in warm storage for two years "
                "after the most recent signed amendment.",
                "Completed contracts move to cold storage after seven years.",
            ),
        ),
        Section(
            "Legal holds",
            (
                "A signed legal hold suspends scheduled deletion. Records remain "
                "preserved until counsel releases the hold and confirms that "
                "disposal may resume.",
            ),
        ),
        Section(
            "Review",
            (
                "The records team reviews the schedule each quarter and records "
                "approved exceptions.",
            ),
        ),
    ),
)

INCIDENT = DocumentSpec(
    filename="incident-response.pdf",
    title="Security Incident Response",
    subtitle="Sample runbook - security operations",
    subject="Fictional security incident response runbook",
    keywords="security, incident, containment, credentials, logs",
    purpose=(
        "This fictional runbook provides safe sample evidence for search and "
        "citation demonstrations."
    ),
    sections=(
        Section(
            "Stabilize and preserve",
            (
                "During an incident, preserve relevant logs, record a timestamped "
                "timeline, and assign one coordinator for operational decisions.",
            ),
        ),
        Section(
            "Contain access",
            (
                "After containment, rotate exposed credentials and verify that "
                "unauthorized access has stopped.",
                "Document each follow-up action with an owner and a due date.",
            ),
        ),
        Section(
            "Closeout",
            (
                "The response lead records lessons learned only after evidence has "
                "been preserved and recovery is confirmed.",
            ),
        ),
    ),
)

BACKUPS = DocumentSpec(
    filename="database-backups.docx",
    title="Database Backup and Restore Runbook",
    subtitle="Sample runbook - platform recovery",
    subject="Fictional database backup and recovery runbook",
    keywords="database, backup, restore, recovery, checksums",
    purpose=(
        "This fictional runbook demonstrates searchable recovery procedures "
        "without referring to a real system or organization."
    ),
    sections=(
        Section(
            "Backup schedule",
            (
                "The catalog receives one encrypted full backup each night and "
                "incremental backups every four hours.",
            ),
        ),
        Section(
            "Recovery proof",
            (
                "A quarterly restore drill loads the newest backup into an isolated "
                "database, verifies checksums, and records recovery time.",
            ),
        ),
        Section(
            "Failure handling",
            (
                "A checksum failure opens an incident and prevents the candidate "
                "recovery point from being promoted.",
            ),
        ),
    ),
)


def _require_generator_dependencies() -> None:
    mismatches = []
    for distribution, expected in EXPECTED_DEPENDENCIES.items():
        observed = version(distribution)
        if observed != expected:
            mismatches.append(f"{distribution}=={expected} (found {observed})")
    if mismatches:
        raise SystemExit(
            "deterministic corpus generation requires " + ", ".join(mismatches)
        )


def _write_zip(path: Path, entries: dict[str, bytes]) -> None:
    """Write a stable OOXML ZIP with canonical order and member metadata."""

    with ZipFile(path, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(entries):
            info = ZipInfo(name, date_time=ZIP_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(
                info,
                entries[name],
                compress_type=ZIP_DEFLATED,
                compresslevel=9,
            )


def _remove_unique_xml_element(
    payload: bytes,
    *,
    tag: bytes,
    attribute: bytes,
    value: bytes,
) -> bytes:
    """Remove one generated self-closing XML element selected by an attribute."""

    pattern = re.compile(
        rb"<"
        + re.escape(tag)
        + rb"\b(?=[^>]*\b"
        + re.escape(attribute)
        + rb'="'
        + re.escape(value)
        + rb'")[^>]*/>'
    )
    sanitized, count = pattern.subn(b"", payload)
    if count != 1:
        raise RuntimeError(
            f"expected one {tag.decode('ascii')} element for "
            f"{attribute.decode('ascii')}={value.decode('ascii')!r}; found {count}"
        )
    return sanitized


def _strip_template_revision_data(payload: bytes) -> bytes:
    """Remove revision identifiers inherited from python-docx's blank template."""

    payload = _RSID_ATTRIBUTE.sub(b"", payload)
    payload = _RSID_LIST.sub(b"", payload)
    payload = _RSID_ELEMENT.sub(b"", payload)
    payload = _WORD_2010_DOCUMENT_ID.sub(b"", payload)
    return _SAVE_PREVIEW_PICTURE.sub(b"", payload)


def _normalize_docx_package(source: Path, destination: Path) -> None:
    with ZipFile(source, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}

    for name in _REMOVED_DOCX_PARTS:
        try:
            del entries[name]
        except KeyError as exc:
            raise RuntimeError(f"expected generated DOCX part {name!r}") from exc

    content_types = entries[_CONTENT_TYPES]
    content_types = _remove_unique_xml_element(
        content_types,
        tag=b"Default",
        attribute=b"Extension",
        value=b"jpeg",
    )
    for part_name in (b"/customXml/itemProps1.xml", b"/docProps/app.xml"):
        content_types = _remove_unique_xml_element(
            content_types,
            tag=b"Override",
            attribute=b"PartName",
            value=part_name,
        )
    entries[_CONTENT_TYPES] = content_types

    package_relationships = entries[_PACKAGE_RELATIONSHIPS]
    for relationship_type in (
        _EXTENDED_PROPERTIES_RELATIONSHIP,
        _THUMBNAIL_RELATIONSHIP,
    ):
        package_relationships = _remove_unique_xml_element(
            package_relationships,
            tag=b"Relationship",
            attribute=b"Type",
            value=relationship_type,
        )
    entries[_PACKAGE_RELATIONSHIPS] = package_relationships
    entries[_DOCUMENT_RELATIONSHIPS] = _remove_unique_xml_element(
        entries[_DOCUMENT_RELATIONSHIPS],
        tag=b"Relationship",
        attribute=b"Type",
        value=_CUSTOM_XML_RELATIONSHIP,
    )

    for name in (
        "word/document.xml",
        "word/settings.xml",
        "word/styles.xml",
        "word/stylesWithEffects.xml",
    ):
        entries[name] = _strip_template_revision_data(entries[name])
    _write_zip(destination, entries)


def _set_style_font(style: object, name: str) -> None:
    style.font.name = name  # type: ignore[attr-defined]
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()  # type: ignore[attr-defined]
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def _set_run_font(run: object, name: str = "Calibri") -> None:
    run.font.name = name  # type: ignore[attr-defined]
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()  # type: ignore[attr-defined]
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def _wrap_pdf_text(text: str, *, font: str, size: float, width: float) -> tuple[str, ...]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if current and stringWidth(candidate, font, size) > width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return tuple(lines)


def generate_pdf(spec: DocumentSpec, path: Path) -> None:
    buffer = BytesIO()
    pdf = canvas.Canvas(
        buffer,
        pagesize=LETTER,
        pageCompression=1,
        invariant=1,
        pdfVersion=(1, 4),
    )
    pdf.setTitle(spec.title)
    pdf.setAuthor(AUTHOR)
    pdf.setSubject(spec.subject)
    pdf.setKeywords(spec.keywords)
    pdf.setCreator(GENERATOR_NAME)
    pdf._doc.info.producer = GENERATOR_NAME

    page_width, page_height = LETTER
    left = inch
    right = page_width - inch
    usable = right - left

    pdf.setFillColor(MUTED)
    pdf.setFont("Helvetica-Bold", 8.5)
    pdf.drawString(left, page_height - 0.72 * inch, "COGNISTORE SAMPLE CORPUS")

    pdf.setFillColor(NAVY)
    pdf.setFont("Helvetica-Bold", 23)
    pdf.drawString(left, page_height - 1.16 * inch, spec.title)
    pdf.setFillColor(MUTED)
    pdf.setFont("Helvetica", 11)
    pdf.drawString(left, page_height - 1.48 * inch, spec.subtitle)

    pdf.setStrokeColor(BLUE)
    pdf.setLineWidth(1.4)
    pdf.line(left, page_height - 1.68 * inch, right, page_height - 1.68 * inch)

    y = page_height - 2.02 * inch
    pdf.setFillColor(LIGHT_BLUE)
    pdf.roundRect(left, y - 0.65 * inch, usable, 0.72 * inch, 6, stroke=0, fill=1)
    pdf.setFillColor(NAVY)
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(left + 0.16 * inch, y - 0.17 * inch, "Fixture purpose")
    pdf.setFillColor(INK)
    pdf.setFont("Helvetica", 9.5)
    purpose_lines = _wrap_pdf_text(
        spec.purpose,
        font="Helvetica",
        size=9.5,
        width=usable - 0.32 * inch,
    )
    for line_index, line in enumerate(purpose_lines):
        pdf.drawString(
            left + 0.16 * inch,
            y - (0.39 + 0.16 * line_index) * inch,
            line,
        )
    y -= 1.02 * inch

    for section in spec.sections:
        pdf.setFillColor(BLUE)
        pdf.setFont("Helvetica-Bold", 12)
        pdf.drawString(left, y, section.heading)
        y -= 0.25 * inch
        pdf.setFillColor(INK)
        pdf.setFont("Helvetica", 10.5)
        for paragraph in section.paragraphs:
            lines = _wrap_pdf_text(
                paragraph,
                font="Helvetica",
                size=10.5,
                width=usable,
            )
            for line in lines:
                pdf.drawString(left, y, line)
                y -= 0.19 * inch
            y -= 0.10 * inch
        y -= 0.13 * inch

    pdf.setStrokeColor(HexColor("#D7DEE5"))
    pdf.setLineWidth(0.6)
    pdf.line(left, 0.64 * inch, right, 0.64 * inch)
    pdf.setFillColor(MUTED)
    pdf.setFont("Helvetica", 8.5)
    pdf.drawString(left, 0.43 * inch, "Project-authored fixture | MIT licensed")
    pdf.drawRightString(right, 0.43 * inch, "1")

    pdf.showPage()
    pdf.save()
    path.write_bytes(buffer.getvalue())


def generate_docx(spec: DocumentSpec, path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x18, 0x20, 0x26)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading = document.styles["Heading 1"]
    _set_style_font(heading, "Calibri")
    heading.font.size = Pt(16)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.line_spacing = 1.1
    heading.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    header_run = header.add_run("COGNISTORE SAMPLE CORPUS")
    _set_run_font(header_run)
    header_run.font.size = Pt(8.5)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor(0x5F, 0x6B, 0x76)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    footer_run = footer.add_run("Project-authored fixture | MIT licensed | 1")
    _set_run_font(footer_run)
    footer_run.font.size = Pt(8.5)
    footer_run.font.color.rgb = RGBColor(0x5F, 0x6B, 0x76)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run(spec.title)
    _set_run_font(title_run)
    title_run.font.size = Pt(23)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x18, 0x3B, 0x56)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.keep_with_next = True
    subtitle_run = subtitle.add_run(spec.subtitle)
    _set_run_font(subtitle_run)
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor(0x5F, 0x6B, 0x76)

    purpose_heading = document.add_paragraph(style="Heading 1")
    purpose_heading.add_run("Fixture purpose")
    purpose = document.add_paragraph(spec.purpose)
    purpose.paragraph_format.keep_together = True

    for item in spec.sections:
        item_heading = document.add_paragraph(style="Heading 1")
        item_heading.add_run(item.heading)
        for text in item.paragraphs:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.keep_together = True

    properties = document.core_properties
    properties.title = spec.title
    properties.author = AUTHOR
    properties.subject = spec.subject
    properties.keywords = spec.keywords
    properties.category = "CogniStore Sample Corpus"
    properties.comments = "Project-authored, fictional, sanitized sample fixture."
    properties.content_status = "Final"
    properties.identifier = "cognistore-content-search-database-backups-v1"
    properties.language = "en-US"
    properties.last_modified_by = GENERATOR_NAME
    properties.created = FIXED_TIMESTAMP
    properties.modified = FIXED_TIMESTAMP
    properties.revision = 1
    properties.version = "1.0"

    with TemporaryDirectory(prefix="cognistore-sample-docx-") as temp_dir:
        intermediate = Path(temp_dir) / spec.filename
        document.save(str(intermediate))
        _normalize_docx_package(intermediate, path)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _manifest(output_dir: Path) -> dict[str, object]:
    retention_sha = _sha256(output_dir / RETENTION.filename)
    incident_sha = _sha256(output_dir / INCIDENT.filename)
    backups_sha = _sha256(output_dir / BACKUPS.filename)
    return {
        "schema_version": SCHEMA_VERSION,
        "name": CORPUS_NAME,
        "license": LICENSE_NAME,
        "bucket": CORPUS_BUCKET,
        "objects": [
            {
                "source": RETENTION.filename,
                "key": "policies/records-retention.pdf",
                "tier": "warm",
                "media_type": "application/pdf",
                "sha256": retention_sha,
            },
            {
                "source": INCIDENT.filename,
                "key": "security/incident-response.pdf",
                "tier": "hot",
                "media_type": "application/pdf",
                "sha256": incident_sha,
            },
            {
                "source": BACKUPS.filename,
                "key": "runbooks/database-backups.docx",
                "tier": "hot",
                "media_type": (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                "sha256": backups_sha,
            },
            {
                "source": BACKUPS.filename,
                "key": "z-archive/database-backups-copy.docx",
                "tier": "cold",
                "media_type": (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                "sha256": backups_sha,
            },
        ],
        "queries": {
            "keyword": {
                "text": "quarterly restore drill",
                "filters": {"bucket": CORPUS_BUCKET},
                "expected_keys": [
                    "runbooks/database-backups.docx",
                    "z-archive/database-backups-copy.docx",
                ],
            },
            "vector": {
                "text": "How does an isolated restore prove backup recovery?",
                "filters": {"bucket": CORPUS_BUCKET, "tier": "hot"},
                "expected_keys": ["runbooks/database-backups.docx"],
            },
            "ask": {
                "text": "What prevents contract deletion during a legal hold?",
                "filters": {
                    "bucket": CORPUS_BUCKET,
                    "mime": "application/pdf",
                },
                "expected_keys": ["policies/records-retention.pdf"],
            },
        },
    }


def _write_manifest(output_dir: Path) -> None:
    serialized = json.dumps(
        _manifest(output_dir),
        ensure_ascii=True,
        indent=2,
        sort_keys=True,
    )
    (output_dir / "manifest.json").write_bytes((serialized + "\n").encode("utf-8"))


def _write_checksums(output_dir: Path) -> None:
    names = sorted((RETENTION.filename, INCIDENT.filename, BACKUPS.filename, "manifest.json"))
    lines = [f"{_sha256(output_dir / name)}  {name}" for name in names]
    (output_dir / "SHA256SUMS").write_bytes(("\n".join(lines) + "\n").encode("ascii"))


def generate_all(output_dir: Path) -> None:
    _require_generator_dependencies()
    output_dir.mkdir(parents=True, exist_ok=True)
    generate_pdf(RETENTION, output_dir / RETENTION.filename)
    generate_pdf(INCIDENT, output_dir / INCIDENT.filename)
    generate_docx(BACKUPS, output_dir / BACKUPS.filename)
    _write_manifest(output_dir)
    _write_checksums(output_dir)


def check(expected_dir: Path) -> None:
    with TemporaryDirectory(prefix="cognistore-content-search-check-") as temp_dir:
        generated_dir = Path(temp_dir)
        generate_all(generated_dir)
        differences = [
            name
            for name in GENERATED_FILES
            if not (expected_dir / name).is_file()
            or not filecmp.cmp(expected_dir / name, generated_dir / name, shallow=False)
        ]
    if differences:
        raise SystemExit("generated corpus differs: " + ", ".join(differences))
    print("content-search sample corpus is deterministic and current")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="generated output directory (default: this asset directory)",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="regenerate in a temporary directory and compare with checked assets",
    )
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    if args.check:
        check(output_dir)
    else:
        generate_all(output_dir)
        print(f"wrote content-search sample corpus to {output_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
